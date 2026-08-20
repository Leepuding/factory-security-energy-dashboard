import csv
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from ..config import Settings


SUPPORTED_TYPES = {"pdf", "docx", "xlsx", "txt", "csv"}
CHUNK_SIZE = 900


class ParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedChunk:
    sequence: int
    content: str
    page_number: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    char_start: int | None = None
    char_end: int | None = None


def detect_file_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_TYPES:
        raise ParseError("仅支持 PDF、DOCX、XLSX、TXT、CSV 文件")
    return suffix


def split_text(text: str, sequence_start: int, **location: int | str | None) -> Iterator[ParsedChunk]:
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        return
    for offset in range(0, len(normalized), CHUNK_SIZE):
        yield ParsedChunk(
            sequence=sequence_start + offset // CHUNK_SIZE,
            content=normalized[offset:offset + CHUNK_SIZE],
            char_start=offset,
            char_end=min(offset + CHUNK_SIZE, len(normalized)),
            **location,
        )


def parse_document(path: Path, file_type: str, settings: Settings, on_progress: Callable[[int, str], None] | None = None) -> list[ParsedChunk]:
    parser = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "xlsx": _parse_xlsx,
        "txt": _parse_txt,
        "csv": _parse_csv,
    }[file_type]
    return list(parser(path, settings, on_progress or (lambda _progress, _message: None)))


def _parse_pdf(path: Path, settings: Settings, on_progress: Callable[[int, str], None]) -> Iterator[ParsedChunk]:
    with fitz.open(path) as pdf:
        if pdf.page_count > settings.max_pdf_pages:
            raise ParseError(f"PDF 页数超过 {settings.max_pdf_pages} 页限制")
        sequence = 0
        for index, page in enumerate(pdf):
            for chunk in split_text(page.get_text("text"), sequence, page_number=index + 1):
                yield chunk
                sequence = chunk.sequence + 1
            on_progress(int((index + 1) / max(pdf.page_count, 1) * 92), f"正在解析第 {index + 1}/{pdf.page_count} 页")


def _parse_docx(path: Path, _settings: Settings, on_progress: Callable[[int, str], None]) -> Iterator[ParsedChunk]:
    document = WordDocument(path)
    sequence = 0
    section_title = "Word 正文"
    paragraph_buffer: list[str] = []

    def flush_paragraphs() -> Iterator[ParsedChunk]:
        nonlocal sequence
        if not paragraph_buffer:
            return
        text = "\n".join(paragraph_buffer)
        for chunk in split_text(text, sequence, sheet_name=section_title):
            yield chunk
            sequence = chunk.sequence + 1
        paragraph_buffer.clear()

    blocks = list(_iter_docx_blocks(document))
    total = max(len(blocks), 1)
    for index, block in enumerate(blocks, start=1):
        if isinstance(block, Paragraph):
            content = block.text.strip()
            if content:
                if _is_section_heading(content):
                    yield from flush_paragraphs()
                    section_title = content
                paragraph_buffer.append(content)
        else:
            yield from flush_paragraphs()
            for chunk in _table_chunks(block, section_title, sequence):
                yield chunk
                sequence = chunk.sequence + 1
        on_progress(10 + int(index / total * 82), "正在提取 Word 段落和表格")
    yield from flush_paragraphs()
    on_progress(92, "Word 文档解析完成")


def _iter_docx_blocks(document: WordDocument) -> Iterator[Paragraph | Table]:
    """Yield paragraph and table blocks in body order so tables keep nearby headings."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _is_section_heading(text: str) -> bool:
    return bool(re.match(r"^[一二三四五六七八九十百]+、", text)) or bool(re.match(r"^\d+[.、]", text))


def _table_chunks(table: Table, section_title: str, sequence_start: int) -> Iterator[ParsedChunk]:
    rows = [[cell.text.strip().replace("\n", " / ") for cell in row.cells] for row in table.rows]
    rows = [[value for value in row] for row in rows if any(row)]
    if not rows:
        return
    header = rows[0]
    label = f"{section_title} · 表格"
    prefix = f"章节：{section_title}\n表头：{' | '.join(header)}\n"
    batch: list[str] = []
    batch_start = 2
    current_size = len(prefix)
    sequence = sequence_start
    for row_number, row in enumerate(rows[1:], start=2):
        line = f"第 {row_number} 行：{' | '.join(row)}"
        if batch and current_size + len(line) + 1 > CHUNK_SIZE:
            yield ParsedChunk(sequence, prefix + "\n".join(batch), sheet_name=label, row_start=batch_start, row_end=row_number - 1)
            sequence += 1
            batch, batch_start, current_size = [], row_number, len(prefix)
        batch.append(line)
        current_size += len(line) + 1
    if batch:
        yield ParsedChunk(sequence, prefix + "\n".join(batch), sheet_name=label, row_start=batch_start, row_end=len(rows))
    elif len(rows) == 1:
        yield ParsedChunk(sequence, prefix, sheet_name=label, row_start=1, row_end=1)


def _parse_xlsx(path: Path, settings: Settings, on_progress: Callable[[int, str], None]) -> Iterator[ParsedChunk]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sequence = 0
    effective_rows = 0
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            batch: list[str] = []
            batch_start: int | None = None
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if not values:
                    continue
                effective_rows += 1
                if effective_rows > settings.max_excel_rows:
                    raise ParseError(f"Excel 有效数据超过 {settings.max_excel_rows} 行限制")
                if batch_start is None:
                    batch_start = row_index
                batch.append(" | ".join(values))
                if len(batch) >= 80:
                    yield ParsedChunk(sequence, "\n".join(batch), sheet_name=worksheet.title, row_start=batch_start, row_end=row_index)
                    sequence += 1
                    batch, batch_start = [], None
            if batch:
                yield ParsedChunk(sequence, "\n".join(batch), sheet_name=worksheet.title, row_start=batch_start, row_end=worksheet.max_row)
                sequence += 1
            on_progress(int((sheet_index + 1) / max(len(workbook.worksheets), 1) * 92), f"正在解析工作表：{worksheet.title}")
    finally:
        workbook.close()


def _parse_txt(path: Path, _settings: Settings, on_progress: Callable[[int, str], None]) -> Iterator[ParsedChunk]:
    text = path.read_text(encoding="utf-8", errors="replace")
    on_progress(55, "正在读取文本文件")
    yield from split_text(text, 0)
    on_progress(92, "文本解析完成")


def _parse_csv(path: Path, settings: Settings, on_progress: Callable[[int, str], None]) -> Iterator[ParsedChunk]:
    sequence = 0
    effective_rows = 0
    batch: list[str] = []
    batch_start: int | None = None
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        for row_index, row in enumerate(csv.reader(source), start=1):
            values = [value.strip() for value in row if value.strip()]
            if not values:
                continue
            effective_rows += 1
            if effective_rows > settings.max_excel_rows:
                raise ParseError(f"CSV 有效数据超过 {settings.max_excel_rows} 行限制")
            if batch_start is None:
                batch_start = row_index
            batch.append(" | ".join(values))
            if len(batch) >= 80:
                yield ParsedChunk(sequence, "\n".join(batch), sheet_name="CSV", row_start=batch_start, row_end=row_index)
                sequence += 1
                batch, batch_start = [], None
    if batch:
        yield ParsedChunk(sequence, "\n".join(batch), sheet_name="CSV", row_start=batch_start, row_end=effective_rows)
    on_progress(92, "CSV 解析完成")
