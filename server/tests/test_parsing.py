import tempfile
import unittest
from pathlib import Path

from docx import Document
from app.config import Settings
from app.services.parsing import ParseError, detect_file_type, parse_document


class ParsingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.settings = Settings(max_excel_rows=3)

    def tearDown(self) -> None:
        self.directory.cleanup()

    def test_txt_records_character_locations(self) -> None:
        path = Path(self.directory.name) / "notes.txt"
        path.write_text("第一段资料\n第二段资料", encoding="utf-8")
        chunks = parse_document(path, "txt", self.settings)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].char_start, 0)
        self.assertIn("第一段资料", chunks[0].content)

    def test_csv_records_row_ranges(self) -> None:
        path = Path(self.directory.name) / "metrics.csv"
        path.write_text("月份,产量\n1月,20\n2月,30\n", encoding="utf-8")
        chunks = parse_document(path, "csv", self.settings)
        self.assertEqual(chunks[0].sheet_name, "CSV")
        self.assertEqual(chunks[0].row_start, 1)
        self.assertEqual(chunks[0].row_end, 3)

    def test_rejects_unknown_type(self) -> None:
        with self.assertRaises(ParseError):
            detect_file_type("unsafe.exe")

    def test_docx_keeps_table_headers_rows_and_section(self) -> None:
        path = Path(self.directory.name) / "production.docx"
        document = Document()
        document.add_paragraph("五、月度生产数据")
        table = document.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text = "月份"
        table.rows[0].cells[1].text = "计划产量"
        table.rows[0].cells[2].text = "实际产量"
        for month, plan, actual in (("1月", "205000", "201500"), ("2月", "168000", "160000")):
            cells = table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = month, plan, actual
        document.save(path)
        chunks = parse_document(path, "docx", self.settings)
        table_chunk = next(chunk for chunk in chunks if chunk.sheet_name and "表格" in chunk.sheet_name)
        self.assertIn("五、月度生产数据", table_chunk.content)
        self.assertIn("表头：月份 | 计划产量 | 实际产量", table_chunk.content)
        self.assertIn("1月 | 205000 | 201500", table_chunk.content)
        self.assertEqual(table_chunk.row_start, 2)
